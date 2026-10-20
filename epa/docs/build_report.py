"""Build the academic report from the implemented project and measured outputs."""
from pathlib import Path
import csv, json, textwrap
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
D=Document()
S=D.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    S[name].font.name='DejaVu Sans'
    S[name].font.color.rgb=RGBColor(0,0,0)
S['Normal'].font.size=Pt(10.5)
S['Normal'].paragraph_format.space_after=Pt(7)
S['Normal'].paragraph_format.line_spacing=1.12
S['Title'].font.size=Pt(28)
S['Heading 1'].font.size=Pt(19)
S['Heading 2'].font.size=Pt(13)
S['Heading 3'].font.size=Pt(11)
for name in ['Heading 1','Heading 2','Heading 3']:
    S[name].paragraph_format.keep_with_next=True
    S[name].paragraph_format.space_before=Pt(10)
    S[name].paragraph_format.space_after=Pt(6)
S['Caption'].font.size=Pt(9)
S['Caption'].font.italic=False
S['Caption'].font.bold=False
for el in S.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)

def setup(sec,kind='portrait'):
    wh={'portrait':(8.27,11.69),'landscape':(11.69,8.27),'foldout':(16.54,11.69)}[kind]
    sec.page_width=Inches(wh[0]);sec.page_height=Inches(wh[1])
    sec.orientation=WD_ORIENT.PORTRAIT if kind=='portrait' else WD_ORIENT.LANDSCAPE
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6)
    sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
setup(D.sections[0])
header=D.sections[0].header.paragraphs[0]
header.text='NETZERO HUB   |   NSW EPA CONTRIBUTION'
header.runs[0].font.size=Pt(8)
footer=D.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('CSCI927/427   |   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
current='portrait'

def new(title,kind='portrait'):
    global current
    sec=D.add_section(WD_SECTION.NEW_PAGE);setup(sec,kind);current=kind
    D.add_heading(title,level=1)

def p(s='',style=None):return D.add_paragraph(s,style)
def h(s):return D.add_heading(s,level=2)
def label(k,v):
    x=p();x.add_run(k+'  ').bold=True;x.add_run(v);return x
def eq(s):
    x=p(s);x.paragraph_format.left_indent=Inches(.12)
    for r in x.runs:r.font.size=Pt(9.5)
    return x
def table(headers,rows,widths=None,size=9.5):
    t=D.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths is None:widths=[6.87/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=str(txt)
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row):cells[i].text=str(txt)
    for ri,row in enumerate(t.rows):
        props=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');props.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');props.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);pr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side,n in [('top',80),('bottom',80),('left',90),('right',90)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            border=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
            pr.append(border)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE8F0');pr.append(shade)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1.05
                for r in pp.runs:r.font.size=Pt(size);r.bold=ri==0
    p().paragraph_format.space_after=Pt(0)
    return t
def fig(file,caption,maxh=6.05):
    file=ROOT/file
    w,h=Image.open(file).size
    mw={'portrait':6.85,'landscape':10.25,'foldout':15.1}[current]
    fw=min(mw,maxh*w/h)
    para=p();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(4)
    para.add_run().add_picture(str(file),width=Inches(fw))
    p(caption,'Caption')
def code(s):
    for line in textwrap.dedent(s).strip('\n').splitlines():
        pp=p(line);pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.0
        rr=pp.runs[0] if pp.runs else pp.add_run(' ');rr.font.name='DejaVu Sans Mono';rr.font.size=Pt(8)
    p()

metrics=json.loads((ROOT/'results/analytics_summary.json').read_text())['processes']
effects=json.loads((ROOT/'models/effects.json').read_text())

D.add_heading('NSW EPA services for NetZero Hub',0)
p('Individual contribution for Tasks A to G','Subtitle')
label('Name','Muhammad Nafees Shihab')
label('Subject','CSCI927/427 Service Oriented Software Engineering, Spring 2026')
label('Provider','New South Wales Environment Protection Authority, NSW EPA')
label('Scope','One provider, two business services, three process models and two implemented microservices')
p('My section of NetZero Hub focuses on the NSW EPA. I have selected two related services: environment protection licence management and pollution incident reporting and investigation. The following sections explain how these services work, how they are modelled and connected, and how a selected part of the design can run as a local prototype.')
p('The prototype supports environmental compliance and responsible incident handling within the wider sustainability platform. It does not calculate carbon savings. The group can connect it to other sustainability providers through shared platform interfaces.')
p('The current implementation includes a React frontend and two FastAPI services. The recorded verification completed 19 API tests, seven frontend helper tests, six browser workflows run individually, and a live HTTP composition check. A seeded simulation generated 300 cases and 2,116 events for three processes. Those are project results from synthetic data, not measurements of NSW EPA operations.')
h('Submission scope')
p('This is one member’s provider section for integration into the six-person group project. The complete group must still meet the provider, process and implementation minima in the specification. Part G includes evidence and agreement fields that require the members’ actual review, ratings and signatures.')
label('Student number and group','To be entered by the student and group')
label('Project boundary','Local prototype with simulated cases; connection to the group platform is a separate integration step')

new('Scope and requirement coverage')
p('The assignment requires the group to work across all aspects and integrate its contributions. Dividing providers among members is an organisation choice; it does not remove the shared responsibility for the final platform.')
table(['Part','This contribution','Group integration remaining'],[
 ('A','1 provider; 2 BSRL specifications','At least 4 other providers, each with at least 2 services'),
 ('B','P1 application, P2 variation, P3 incident; English and FOL effects','At least 2 additional realistic processes; combine the catalogue'),
 ('C','SoaML architecture, interfaces, contracts, protocols and composition','Integrate all providers and the platform’s common contracts'),
 ('D','Business, Application and Technology layers for the EPA slice','Connect the other providers and shared platform capabilities'),
 ('E','React EPA module; 2 independent APIs with separate databases','At least 1 more implemented service and a unified demonstration'),
 ('F','3 executed simulations; logs; Alpha and Inductive mining; conformance','Review and integrate the analysis with the final report'),
 ('G','Evidence worksheet and six agreement/signature fields','Actual peer ratings, each member’s agreement and signature')], [.5,3.05,3.32])
p('Part A defines the two services. Part B describes their processes and effects. Parts C and D explain the architecture, Part E presents the implementation, Part F analyses the simulated event logs, and Part G records contribution evidence for group review.')

new('A Service identification and specification')
h('Provider and service catalogue')
p('The selected provider is the NSW Environment Protection Authority (NSW EPA), the state’s independent environmental regulator. Its work includes licensing regulated activities, responding to pollution and encouraging better environmental performance (NSW EPA, 2024d, 2025a). The two selected services give NetZero Hub users a practical way to manage environmental obligations and raise pollution concerns.')
table(['ID','Business service','Consumer and value','Processes'],[
 ('EPA-001','Environment Protection Licence Application and Management','Businesses/applicants manage a proposed activity and its licence conditions','P1 application; P2 variation'),
 ('EPA-002','Pollution Incident Reporting and Investigation','Citizens or businesses report a pollution concern and receive an appropriate outcome','P3 reporting and investigation')],[.7,2.5,2.45,1.22])
p('EPA-001 groups related licence capabilities into one business service. The implemented subset includes new applications and variations. Transfers, surrenders, applicable load-based licensing annual returns and other reporting obligations belong to the wider licence service (NSW EPA, 2026a, 2026b, 2026d). EPA-002 groups intake, triage and investigation into one service. The internal licence-context lookup supports these services and is not counted as a third business service.')
h('BSRL method and contract interpretation')
p('The catalogue uses EPA-001 and EPA-002 consistently. The development UI labels EPA001 and EPA002 refer to those same services. P1, P2 and P3 are local EPA process identifiers, not additional business services.')
p('The following specifications use the Business Service Representation Language (BSRL): service description, goal, preconditions, postconditions, assumptions, inputs and outputs with formats, resources, key steps, and QoS/contractual factors. A goal states intended value; a postcondition states the observable result of execution (Ghose et al., 2011).')
p('All API latency, availability, timer and recovery values below are proposed project targets. They are not promises published by NSW EPA and have not been established by a load test. Real licence fees and statutory decision obligations are not reproduced as fixed numerical promises. The prototype records references only; it has no payment integration.')
h('NetZero Hub relationship')
p('NetZero Hub provides a common entry point for the group’s renewable energy, waste, EV charging, transport and environmental services. The EPA pages let applicants and officers submit requests and inspect progress through the same portal. The Hub consumes the EPA service contracts while providing the user-facing experience. LicenceService and IncidentService are project components representing the NSW EPA domain; no official EPA API integration is claimed.')

new('A BSRL specification for EPA-001')
label('Service description','EPA-001 — Environment Protection Licence Application and Management')
label('Category and provider','Environmental regulation and licensing; NSW EPA')
label('Goal','Help eligible applicants seek an Environment Protection Licence and help licence holders manage changes, conditions and reporting obligations. Within NetZero Hub, the selected workflow makes each request and its outcome traceable.')
label('Preconditions','The applicant provides identifiable premises, activity and organisation details and has authority to act for the applicant organisation. The activity must be within the relevant licensing framework. Required documents, applicable fees and consent information are checked before assessment. A management request identifies an existing licence (NSW EPA, 2026b, 2026c, 2026d).')
label('Postconditions','The request and its outcome are recorded. An approved application results in a licence with conditions; an approved variation updates the relevant conditions. A request for further information or a refusal is communicated. Transfer, surrender and reporting requests have their own outcomes in the wider service. Approval is conditional, so it is not promised for every submission (NSW EPA, 2026a, 2026b, 2026c, 2026d).')
label('Assumptions','The applicant supplies truthful information and can obtain any required external approvals. The responsible officers have access to relevant evidence and decision authority. External approval and payment checks are represented by references in this prototype; their validity is not independently verified.')
table(['Business input','Representation in the selected prototype'],[
 ('Applicant, premises and activity details','Validated JSON strings'),
 ('Supporting documents, consent and fee evidence','References only; document IDs as strings, consent/fee as optional strings'),
 ('Existing licence and proposed change','Licence ID, conditions list and rationale'),
 ('Assessment and decision information','Officer-entered results, reason and expected case version')],[3.2,3.67])
table(['Business output','Representation in the selected prototype'],[
 ('Acknowledgement, reference and status','JSON case ID, process, owner, state, version and data'),
 ('Licence or approved variation','Prototype licence ID, active flag, version and conditions'),
 ('Decision and progress history','Authorised case response and completed-task audit events')],[3.2,3.67])

new('A EPA-001 delivery and implementation scope')
label('Resources','Applicant, licensing officer, environmental assessor, decision officer, case documents, licence registry, LicenceService and its owned database.')
label('Key steps','P1: receive → check → request/recheck if needed → two assessments → decide → issue/refuse → publish. P2: receive → verify → assess change → decide → update/refuse → publish.')
label('Quality of service','The business service needs accurate records, traceable decisions, appropriate access control and clear communication. The Charter supports fair treatment, risk-sensitive responses and timeframes suited to the matter (NSW EPA, 2025a). The proposed technical constraints below use the BSRL <factor, range> form (Ghose et al., 2011).')
table(['QoS pair for the Hub prototype','Status'],[
 ('<response time at p95, at most 2 seconds>','Design target; an agreed load test is still required'),
 ('<monthly availability, at least 99.5 percent>','Design target; not measured by local tests'),
 ('<duplicate creation for an identical idempotent request, zero>','Covered by automated behaviour tests'),
 ('<case and licence update integrity, atomic local commit>','Implemented with a database transaction and version checks')],[4.5,2.37])
label('Delivery schedule','<receipt, on successful submission>; <updated outcome, after the decision commits>. Actual EPA assessment timing depends on the applicable process and information required (NSW EPA, 2025a, 2026c). The prototype’s 30-day information window is an explicit modelling assumption, not an EPA deadline.')
label('Payment schedule','Applicable official licence fees follow the EPA’s current instructions. Annual administrative fees and other charges may apply to real licences (NSW EPA, 2026b). The prototype accepts a fee reference and collects no money.')
label('Penalty or remedy','The official guidance does not establish a customer compensation schedule for delays. No discount, refund or fixed contractual penalty is assumed. A prototype failure is recorded for correction; regulatory consequences for a licence holder’s non-compliance are a separate matter (NSW EPA, 2024c, 2026b).')
h('Selected executable scope')
p('The model implements standard-case application and variation paths. An authenticated demo applicant owns each request. Approval creates one prototype licence or updates its conditions and version. Refusal, invalidity and incomplete closure preserve a reasoned outcome. Annual reporting, transfer, surrender, payments and official document exchange remain outside the implemented subset.')

new('A BSRL specification for EPA-002')
label('Service description','EPA-002 — Pollution Incident Reporting and Investigation')
label('Category and provider','Environmental incident management; NSW EPA')
label('Goal','Record a pollution concern, determine the appropriate handling route and preserve a reasoned investigation outcome with supporting evidence.')
label('Preconditions','A pollution concern is observed or suspected and enough information is provided to identify the issue. Responsibility is checked before an EPA investigation proceeds. Immediate threats require the appropriate emergency channel. The official service accepts anonymous information (NSW EPA, 2025a, 2025c); authentication is only a condition of this Hub demonstration.')
label('Postconditions','The concern is recorded and its handling route is identified. The reporter receives guidance, a referral direction or an investigation outcome as appropriate. An investigation may collect evidence and lead to a reasoned regulatory response. A report alone does not establish an offence or guarantee prosecution (NSW EPA, 2024c, 2025b, 2025c).')
label('Assumptions','Information and evidence are provided in good faith. Officers can obtain the records and access needed for their work where permitted. Risk, jurisdiction and investigative decisions require human judgement. Referral of a person’s details would need the applicable consent or other authority (NSW EPA, 2025a); the prototype sends no referral.')
table(['Business input','Representation in the selected prototype'],[
 ('Issue, location, date and time, description and impact','Category/location/description strings and timezone-aware timestamp'),
 ('Photographs, recordings, documents or witness information','Officer-entered evidence summary; binary attachments are not stored'),
 ('Optional premises or licence context','Prototype licence ID or no reference'),
 ('Risk, jurisdiction and assessment information','Officer commands with values and expected case version')],[3.3,3.57])
table(['Business output','Representation in the selected prototype'],[
 ('Receipt, reference and guidance','JSON case and immediate guidance text'),
 ('Progress and investigation outcome','Authorised case status and recorded conclusion'),
 ('Supporting record','Evidence/context summaries and audit events')],[3.3,3.57])

new('A EPA-002 delivery and implementation scope')
label('Resources','Reporter, intake officer, investigator, decision role, evidence records, IncidentService, its own database, and the LicenceContext interface of EPA-001.')
label('Key steps','Register → triage → guidance/referral or assignment → collect evidence and review context → assess evidence → record action/no action → close → publish.')
label('Quality of service','Reports should be recorded accurately, sensitive information should be handled appropriately, and urgency should reflect the risk to people and the environment. Decision reasons and progress should be understandable to the recipient. Confidentiality is subject to lawful disclosure requirements (NSW EPA, 2025a).')
table(['QoS pair for the Hub prototype','Status'],[
 ('<receipt response time at p95, at most 2 seconds>','Design target; not benchmarked'),
 ('<access checks, every protected request>','Implemented; selected owner/role cases tested'),
 ('<licence lookup timeout, 2 seconds>','Configured HTTP timeout'),
 ('<case loss on lookup failure, zero committed state changes>','503 response; case version and stored state are preserved')],[4.5,2.37])
label('Delivery schedule','<receipt and safety guidance, successful submission response>; <outcome access, after the relevant decision commits>. Official urgent reporting uses Environment Line, available 24/7, or emergency services where needed. Investigation time varies with the evidence and complexity; the prototype sets no statutory completion promise (NSW EPA, 2025a, 2025c).')
label('Payment schedule','No reporting charge is modelled and the prototype takes no payment. A polluter’s possible regulatory liabilities are outside the reporter’s service-payment schedule.')
label('Penalty or remedy','No commercial compensation schedule is specified. A service error should be recorded and corrected. Possible warnings, notices or prosecution concern regulatory action against a responsible party and are not penalties imposed on the reporter for using the service. Decisions depend on evidence and the relevant circumstances (NSW EPA, 2024c, 2025b).')
h('Selected executable scope')
p('The prototype uses named demo accounts so that case ownership can be demonstrated. It does not implement an anonymous reporting channel. Triage selects emergency guidance, referral guidance or an EPA investigation. Guidance ends that prototype case; any later real emergency follow-up is outside this workflow.')
p('An investigation requires both an evidence summary and a licence-context review. The context can explicitly record that no licence was supplied or found. A failed service lookup remains pending for retry. “Action recorded” is an officer-entered outcome in the project database; it does not issue a notice or commence a prosecution.')

new('B Process basis and modelling decisions')
p('The models use publicly documented service activities as their starting point. The detailed task allocation, sequencing, digital messages and branch rules are proposed project design. This distinction is necessary because public guidance does not disclose the EPA’s internal process engine or complete decision procedure.')
table(['Source evidence','BPMN mapping','Design added by this project'],[
 ('Application details, completeness and assessment criteria (NSW EPA, 2026a, 2026c)','P1 A1–A4 intake; A6 environmental review; A7 suitability/consent; A8 decision','Separate parallel assessments; a 30-day information timer; prototype issue/refusal and outcome states'),
 ('Licence variation and changes to conditions (NSW EPA, 2026c, 2026d)','P2 V1 request; V4 change assessment; V5 decision; V6 conditions update','Ownership/version eligibility, an invalid branch, optimistic concurrency and digital polling'),
 ('Pollution reporting, investigation, routing and confidentiality (NSW EPA, 2025c, 2026f)','P3 I1 intake; I2 triage; I6 evidence; I8 assessment; I9/I10 outcome','Explicit emergency/referral terminal states; parallel licence-context review; separate service API')],[2.55,2.05,2.27])
h('BPMN elements and responsibilities')
p('A black-box consumer pool represents the applicant or reporter interacting through the Hub and communicates with a provider pool containing intake/assessment and registry/outcome lanes. Sequence flows stay inside the provider pool. The model files use message flows between pools, a message start event, end events, user and service tasks, send/receive tasks, XOR decisions, AND splits and joins, data objects and text annotations. P1 adds an interrupting boundary timer attached to the information-receive task and a rework loop. These choices follow the BPMN event, gateway and pool semantics (Object Management Group, 2014).')
p('P1’s AND join waits for A6 and A7. P2’s XOR paths retain exactly one of invalid, approved or refused. P3’s first XOR handles emergency guidance; the next handles jurisdiction; its AND join requires I6 and I7 before evidence assessment. The diagrams show branch guards; commands enforce the selected prerequisites. The implementation’s timer is evaluated when an officer invokes expire, rather than by a background scheduler.')
p('In the combined report, qualify local identifiers when needed: EPA.P1 is the application process and EPA.P1.A1 is its receive task. The EPA files retain P1 and A1 so their diagrams, transitions and event logs continue to match. Consumer actions are outside the black-box pool’s internal model; the contract sequence diagrams show their messages.')
p('Tasks such as “publish outcome” mean making data available for an authorised GET request. The project does not implement email delivery, automatic emergency calls, actual referrals or regulatory enforcement. Application completeness is intentionally simplified to three reference categories. A complete real-world system would need the current official forms, legislation, exceptions and delegated decision rules.')

new('B EPA guidance and process design')
p('EPA guidance explains the business checks and responsibilities used in the models. The table connects each official source to the relevant process decision. Detailed task sequencing and software behaviour remain project design choices.')
table(['Official EPA source','Supported business detail','Effect on the model'],[
 ('Environment protection licences (NSW EPA, 2026b)','Application/variation channels, licence conditions and compliance monitoring','A1, A6, A8 and V1/V4/V6; conditions are maintained as versioned data'),
 ('Report pollution (NSW EPA, 2025c)','Incident details, emergency routing, responsible authority, evidence and investigation','I1, I2, I3, I4, I6 and I8; investigation is one possible handling route'),
 ('Customer Service Charter (NSW EPA, 2025a)','Anonymous information, qualified confidentiality, risk priority and consent for referrals','Official anonymity is recognised; demo accounts and guidance-only referrals are explicit limits'),
 ('Authorised Officers and Enforcement Officers (NSW EPA, 2020)','Powers depend on appointment and function','Separate officer roles; applicant/reporter identities cannot approve or conclude a case'),
 ('Regulatory Policy (NSW EPA, 2024c)','Proportionate responses depend on evidence, harm, history and circumstances','I8/I9 require an officer conclusion and reason; a boolean is only a simplified routing input'),
 ('Prosecution Guidelines (NSW EPA, 2025b)','Prosecution requires an evidential assessment and discretion','I9 records a selected response; it does not automatically prosecute'),
 ('Public registers (NSW EPA, 2026e)','Selected licence and regulatory information is publicly searchable','Licence context is a useful integration dependency; the demo uses its own minimal registry')],[2.3,2.25,2.32],9.0)
h('Further requirements outside the selected workflow')
p('Financial-capacity assessment may inform licensing and other decisions, so a future version of A7 could request a specialist assessment. The financial-capacity webpage provides a worked example for a fictional company; it is not an EPA case log. Financial assurance is a way to secure environmental obligations; it is not a customer refund or a penalty for a breach (NSW EPA, 2024a, 2024b).')
p('The EPA’s compliance audit program describes checks against licence conditions, findings and follow-up action. This could support a separate proactive inspection process. The current P3 starts with a pollution report, so it does not treat an audit campaign as the same process (NSW EPA, 2022).')

for proc,title,desc in [
 ('P1','B P1 Licence application model','Incomplete cases request information and return for checking, or close through the timer. Complete cases require both assessments before a decision. The shown C1–C3 scenarios use the direct complete path T0; the indexed text generalises them to rework histories.'),
 ('P2','B P2 Licence variation model','Eligibility failure ends with an unchanged licence. Eligible cases pass through change assessment and decision. Approval replaces conditions and increments the licence version. C4–C6 and V1–V8 are indexed below.'),
 ('P3','B P3 Pollution investigation model','Emergency and other-authority cases produce guidance. An EPA investigation requires both evidence and context review before its conclusion. C7–C9 and I1–I12 are indexed below.')]:
    new(title,'foldout')
    fig('diagrams/'+proc+'_BPMN.png',desc,maxh=9.0)
    p('English and FOL annotations are included in the editable BPMN file and this figure. Full-size SVG and indexed text are provided for detailed reading. Task IDs are consistent across models, code and event logs.','Caption')

new('B Semantic annotation method')
p('An immediate effect describes what becomes true after one completed task. A cumulative effect combines the effects and guards along one feasible execution history. Using semantic effect annotation, a scenario is represented as a triple ⟨A, L, X⟩: cumulative assertion A, path label L, and exclusion set X containing incompatible alternative prefixes. XOR alternatives remain separate scenarios; an AND join combines only compatible branch scenarios (Hinge et al., 2009).')
h('Variables and interpretation')
table(['Symbol','Meaning'],[('a, v, i','Application, variation and incident case identifiers'),('l, k','Licence identifier and its current integer version before an update'),('c, e, f, q, s','Boolean completeness, assessment, suitability, eligibility and substantiation results'),('d, r, j','Decision; risk class; jurisdiction class'),('∧, ¬, ∃, ∀, ⇒','And, not, exists, for all, implies'),('BaseA, BaseV, BaseI, C1–C9','Abbreviations expanded in the cumulative scenario pages')],[1.2,5.67])
p('Free variables refer to the case and values in the current scenario; each formula is a schema instantiated for that case. Constants include true, false, approve, refuse, emergency, nonEmergency, epa and other. “state” denotes one current terminal outcome in these annotations. Historical task completion facts are retained separately from mutable result facts.')
h('Background knowledge and accumulation')
eq('∀x∀u∀w ((state(x,u) ∧ state(x,w)) ⇒ u=w)')
eq('∀x∀u∀w ((completeResult(x,u) ∧ completeResult(x,w)) ⇒ u=w)')
eq('∀l∀m∀n ((version(l,m) ∧ version(l,n)) ⇒ m=n)')
eq('∀a ((Application(a) ∧ decided(a,approve)) ⇒ envResult(a,true) ∧ fitResult(a,true))')
eq('∀v ((Variation(v) ∧ decided(v,approve)) ⇒ eligibleResult(v,true) ∧ changeResult(v,true))')
p('Application(a) and Variation(v) identify disjoint case types and are persistent background facts. Variables u, w, m and n range over the relevant result or version values.')
p('The same single-value rule applies to other current result predicates. Use acc(A,e) to add a new effect while replacing older clauses that conflict with it under this background knowledge. Thus A4 removes awaitingInfo, a successful recheck replaces completeResult(false), and V6 replaces version(l,k) with version(l,k+1). Simple conjunction without those replacements would create inconsistent scenarios (Hinge et al., 2009).')
p('The formulas specify the model and are checked by inspection against the guards and transitions. The code enforces the selected workflow rules; it does not implement a general first-order logic engine or automatically derive these scenarios.')

for prefix,title in [('A','B Immediate effects for P1'),('V','B Immediate effects for P2'),('I','B Immediate effects for P3')]:
    new(title)
    p('Each row is the immediate effect of the task carrying the same ID in the BPMN model. English describes the effect; the formula formalises it. Accumulated history and branch guards are added separately.')
    rows=[(key,val['task']+'\n'+val['english'],val['fol']) for key,val in effects.items() if key.startswith(prefix)]
    table(['ID','Task and English effect','First order logic'],rows,[.45,3.17,3.25],9.2)
    if prefix=='A':p('A5 is incomplete-case closure under the project’s timer policy. A9 issues only a prototype record. A6 and A7 may complete in either order; both are required at C1.')
    if prefix=='V':p('V6 uses the pre-update version k and atomically records k+1. V3 and V7 preserve the existing conditions. A stale licence version prevents the approval command from committing.')
    if prefix=='I':p('A “recorded” action or referral is a prototype data record. Context can explicitly report no matching licence or no supplied reference; an unavailable lookup remains incomplete.')

new('B Cumulative scenarios for P1')
p('The diagram’s BaseA(a) abbreviates received(a) ∧ checked(a) ∧ completeResult(a,true). Let Tₖ be A1 followed by k information-response cycles (A2[false], A3, A4), then A2[true]. For k=0 the application is complete immediately. For k>0 add Hₖ = infoRequested(a) ∧ infoReceived(a) ∧ ¬awaitingInfo(a); H₀ is true.')
label('C1 at the assessment join','The application is complete and both assessment results are recorded. Both results remain available regardless of completion order.')
eq('C1ₖ(a) = BaseA(a) ∧ Hₖ ∧ envChecked(a) ∧ envResult(a,e) ∧ fitChecked(a) ∧ fitResult(a,f)')
label('Path label','L(C1ₖ) = Tₖ · (A6 || A7). The symbol || joins two required branches, not an XOR choice.')
label('C2 after A9','A reasoned approval exists and one prototype licence has been issued; its guard requires e=f=true.')
eq('C2ₖ(a) = C1ₖ(a) ∧ decided(a,approve) ∧ reasonRecorded(a) ∧ state(a,approved) ∧ ∃l issued(l,a)')
label('C3 after A10','The assessed application has a reasoned refusal and no licence was issued for this case.')
eq('C3ₖ(a) = C1ₖ(a) ∧ decided(a,refuse) ∧ reasonRecorded(a) ∧ state(a,refused) ∧ ¬∃l issued(l,a)')
label('Waiting scenario after A3','The latest completeness result is false and the case awaits information. Earlier completed rounds may remain as history.')
eq('received(a) ∧ checked(a) ∧ completeResult(a,false) ∧ infoRequested(a) ∧ awaitingInfo(a)')
label('Timer scenario after A5','The incomplete case closes; the current wait ends. After A11 the outcome is available.')
eq('received(a) ∧ checked(a) ∧ completeResult(a,false) ∧ infoRequested(a) ∧ ¬awaitingInfo(a) ∧ state(a,incompleteClosed)')
p('For a timer path after earlier responses, retain infoReceived(a) as a historical fact. At any terminal A11 add outcomeAvailable(a) to that path’s assertion. Never combine approval, refusal and timeout assertions into one conjunction.')
h('Path exclusions and consistency')
p('At every unrolled XOR choice, put the opposite guarded prefix into X. For Tₖ, X therefore contains the complete-prefix alternative at each earlier incomplete visit and the incomplete-prefix alternative at the final complete visit. C2 adds the A8[refuse] prefix; C3 adds A8[approve]. A timeout scenario excludes the response-before-deadline prefix of the same waiting occurrence. Rework occurrences are indexed, so a later complete result does not conflict with an earlier incomplete history.')
p('At C1, A6 and A7 scenarios may be combined only when their Tₖ histories match and neither path is in the other’s exclusion set. Their exclusion sets are unioned. This prevents combining an assessment from one alternative application history with another (Hinge et al., 2009).')

new('B Cumulative scenarios for P2')
p('Let BaseV(v,l) = received(v) ∧ refersTo(v,l) ∧ eligibilityChecked(v) ∧ eligibleResult(v,true), with version(l,k) in the incoming licence state. Let U = V1 · V2[true] · V4. The invalid prefix is V1 · V2[false].')
label('C4 after V4','The licence is eligible and the proposed change has been assessed; e records the assessment result.')
eq('C4(v,l) = BaseV(v,l) ∧ changeAssessed(v) ∧ changeResult(v,e)')
label('Scenario','⟨C4(v,l), U, {V1 · V2[false]}⟩')
label('C5 after V6','The approved variation has a reasoned decision, updated conditions and the next licence version. Its guard requires e=true and an unchanged pre-update version.')
eq('C5(v,l) = C4(v,l) ∧ decided(v,approve) ∧ reasonRecorded(v) ∧ state(v,approved) ∧ varied(l,v) ∧ version(l,k+1)')
label('Scenario','L = U · V5[approve] · V6; X = {V1 · V2[false], U · V5[refuse]}. Accumulation removes version(l,k).')
label('C6 after V7','The assessed variation is refused, with a reason recorded and the existing licence conditions preserved.')
eq('C6(v,l) = C4(v,l) ∧ decided(v,refuse) ∧ reasonRecorded(v) ∧ state(v,refused) ∧ unchanged(l,v)')
label('Scenario','L = U · V5[refuse] · V7; X = {V1 · V2[false], U · V5[approve]}. The incoming version(l,k) persists.')
label('Invalid scenario after V3','Verification failed; no licence conditions change.')
eq('received(v) ∧ refersTo(v,l) ∧ eligibilityChecked(v) ∧ eligibleResult(v,false) ∧ state(v,invalid) ∧ unchanged(l,v)')
label('Scenario','L = V1 · V2[false] · V3; X = {V1 · V2[true]}. The incoming licence version persists.')
p('At V8 add outcomeAvailable(v) separately to the approved, refused or invalid scenario. The XOR merge does not require all branches to finish and does not make all their effects true together.')
h('State and concurrency interpretation')
p('The case version and licence version are different. Each successful command advances the case version; an approved variation alone advances the licence version. A variation can become stale after another variation changes the same licence. The approval guard checks the current licence version inside the transaction, so the stale request cannot overwrite newer conditions.')
p('unchanged(l,v) is relative to this request: it means that variation v did not alter l. It does not assert that no other case has ever changed the licence. This scope avoids an unnecessarily strong global claim.')

new('B Cumulative scenarios for P3')
p('BaseI(i) = received(i) ∧ triaged(i) ∧ risk(i,nonEmergency) ∧ jurisdiction(i,epa) ∧ assigned(i). Let N = I1 · I2[nonEmergency,epa] · I5; E = I1 · I2[emergency]; R = I1 · I2[nonEmergency,other].')
label('C7 at the investigation join','An EPA investigation has both evidence and licence context recorded.')
eq('C7(i) = BaseI(i) ∧ evidenceRecorded(i) ∧ contextRecorded(i)')
label('Scenario','L = N · (I6 || I7); X = {E, R}. Both branches carry the same intake history and compatible exclusions.')
label('C8 after I11 on the action branch','Evidence was evaluated as substantiated, an action was recorded and the investigation closed.')
eq('C8(i) = C7(i) ∧ evaluated(i) ∧ substantiatedResult(i,true) ∧ actionRecorded(i) ∧ state(i,actionRecorded) ∧ investigationClosed(i)')
label('Scenario','L = N · (I6 || I7) · I8[true] · I9 · I11; X = {E, R, N · (I6 || I7) · I8[false]}.')
label('No further action scenario after I11','Both prerequisites are complete, the concern is not substantiated for this prototype decision, and the investigation closes with a reasoned no-action record.')
eq('C7(i) ∧ evaluated(i) ∧ substantiatedResult(i,false) ∧ noActionRecorded(i) ∧ state(i,noFurtherAction) ∧ investigationClosed(i)')
label('Scenario','L = N · (I6 || I7) · I8[false] · I10 · I11; exclude E, R and the corresponding I8[true] prefix.')
label('C9 after I4','The non-emergency report falls outside the model’s EPA jurisdiction, so referral guidance is available.')
eq('C9(i) = received(i) ∧ triaged(i) ∧ risk(i,nonEmergency) ∧ jurisdiction(i,other) ∧ guidance(i,referral) ∧ state(i,referred)')
label('Scenario','L = R · I4; X = {E, I1 · I2[nonEmergency,epa]}.')
label('Emergency scenario after I3','Immediate danger is identified and emergency guidance is available.')
eq('received(i) ∧ triaged(i) ∧ risk(i,emergency) ∧ jurisdiction(i,j) ∧ guidance(i,emergency) ∧ state(i,emergencyGuidance)')
label('Scenario','L = E · I3; X contains I1 · I2[nonEmergency]. At I12 add outcomeAvailable(i) to each terminal scenario separately.')
p('An emergency or referral scenario has no evidence/context join and no investigationClosed assertion. Recording guidance is not evidence that a call, referral or enforcement action was carried out.')

new('C Service architecture and roles','landscape')
fig('diagrams/C1_SoaML_architecture.png','Figure C1. Project collaboration structure for the two business services, using SoaML notation (Object Management Group, 2012).',5.55)
p('NetZeroHub binds the consumer roles to request ports; the provider participant exposes service ports. LicenceManagementContract and PollutionManagementContract specify the obligations of the two roles. The platform aggregates access while the provider decision role supplies the business decisions. The Hub is the consumer at this provider boundary and the provider of the shared user experience. These are proposed relationships, not an assertion about EPA internal software.','Caption')

new('C Service interfaces and data','landscape')
fig('diagrams/C2_SoaML_interfaces.png','Figure C2. ServiceInterface and MessageType views. Method names are logical design names; the HTTP binding is specified below.',5.9)

new('C Contracts and HTTP binding')
p('A service contract describes both parties and their permitted interaction (Object Management Group, 2012). The provider promises a receipt for a valid accepted request and an authorised status response; the consumer supplies accurate structured data, an idempotency key for creation and a current case version for mutation. The implementation uses JSON request/reply and polling. There is no callback delivery obligation.')
table(['Contract','Preconditions and obligations','Postconditions and faults'],[
 ('Licence management contract','Consumer owns the request/licence; officer supplies assessments and a reasoned decision; all required checks complete','Stable case; one issue or versioned variation on approval. 403 for ownership/role failure; 409 for stale state or missing prerequisite'),
 ('Pollution management contract','Consumer supplies observation details; officer records triage; evidence and context precede conclusion','Guidance or recorded investigation outcome. 422 for malformed input; 503 leaves context review pending on lookup failure'),
 ('Licence context contract','Internal or officer token; opaque prototype licence ID','Returns id, active, version and conditions, without owner details. 404 means no match; other lookup failures are retriable')],[1.55,2.65,2.67])
table(['Logical operation','HTTP binding','BPMN tasks'],[
 ('submitApplication','POST /applications','A1'),('submitVariation','POST /variations','V1'),('submitReport','POST /incidents','I1'),
 ('getCase','GET /cases/{case_id} on the owning service','Outcome access after A11/V8/I12'),
 ('provideInformation','POST /cases/{id}/commands; action=supply','A4'),
 ('check and assess','Same command endpoint; check, assess_environment, assess_suitability','A2/A3, A6, A7'),
 ('decide or vary','Same command endpoint; verify, assess, decide','A8–A11 or V2–V8'),
 ('triage and investigate','Same command endpoint; triage, collect_evidence, review_licence, conclude','I2–I12'),
 ('getContext','GET /internal/licences/{id}/context','Supports I7')],[1.45,3.35,2.07],9.1)
p('Message types are bound to Pydantic schemas in the source and to the exported OpenAPI documents. A CaseReceipt is the returned case object; caseId/processId in the design map to id/process in JSON. InformationCommand carries expected_version, action=supply and updated values. An officer-only command interface complements the public request interfaces.')

new('C Licence contract behaviour','landscape')
fig('diagrams/C3_Licence_contract.png','Figure C3. Application protocol. Information supply repeats only while required; assessment and approval/refusal require a complete case.',5.6)
p('For P1, timeout leads to A5 and a terminal outcome instead of assessment. For P2, the contract uses submitVariation → verify → assess → decide → getCase. An invalid verification returns a terminal outcome directly; approval increments the licence version. P2 has one change assessment and no information loop.','Caption')

new('C Incident contract behaviour','landscape')
fig('diagrams/C4_Incident_contract.png','Figure C4. Mutually exclusive guarded interaction fragments for guidance and investigation. Officer actions are routed through the Hub role.',5.6)
p('Evidence collection can precede or follow the context review. The figure shows one permitted order. A context timeout interrupts the investigation fragment before conclusion; retry is explicit and the case stays saved. All paths eventually expose their permitted outcome through polling.','Caption')

new('C Composition and interoperability','landscape')
fig('diagrams/C5_SoaML_composition.png','Figure C5. Project participant decomposition and the LicenceContext dependency, using SoaML notation (Object Management Group, 2012).',5.4)
p('IncidentService orchestrates a read-only call to LicenceService when an investigator requests context review. Each service owns its data; the incident service receives a minimal context snapshot with a licence version. The composition joins that review with evidence before conclusion. Shared JSON types, opaque identifiers, explicit versions and timezone-aware timestamps support interoperability. No cross-database transaction or event broker is needed for this selected read-only composition.','Caption')

new('C Contract effects and platform integration')
p('The EPA contract messages achieve the process effects defined in Part B. The table gives the main guarantees at the contract boundary; the complete cumulative scenarios also retain their earlier task history and branch conditions. A successful HTTP response does not itself establish a regulatory approval.')
table(['Interaction and BPMN tasks','Required state','Guaranteed effect in this prototype'],[
 ('Approve an application: A6–A9, A11','Environmental and suitability results both pass; the officer supplies a reason','state(a,approved) ∧ ∃l issued(l,a) ∧ outcomeAvailable(a)'),
 ('Approve a variation: V2, V4–V6, V8','Eligibility and change assessment pass; the licence version remains current','state(v,approved) ∧ varied(l,v) ∧ version(l,k+1) ∧ outcomeAvailable(v)'),
 ('Conclude an investigation: I6–I12','Evidence and licence-context review are complete; an officer records the conclusion','investigationClosed(i) ∧ outcomeAvailable(i), with either the action or no-action outcome'),
 ('Provide guidance: I2–I4, I12','Triage selects an emergency or other-authority route','The relevant guidance and outcome are recorded; no investigation-closure effect is asserted'),
 ('Retry a context lookup: I7','The previous lookup failed before context was committed','A 503 leaves the case unchanged. A later successful review records contextRecorded(i)')],[1.95,2.25,2.67],9.1)
h('Relationship to the other provider modules')
p('Each provider keeps responsibility for its own business rules and data. The shared portal supplies navigation and a common identity, then calls the appropriate domain API. Solar installations, waste bookings and charging sessions do not automatically create EPA applications. Any future cross-provider environmental check needs an explicit contract and a justified business rule.')
p('EPA status is obtained through authorised GET requests. A waste-collection module may use a completion callback without requiring the EPA module to implement the same mechanism. The common platform must preserve the delivery obligations of each contract. EPA “publish outcome” means that the owner can retrieve it; it does not mean an email or callback was sent.')
h('Composition boundary')
p('The implemented composition is IncidentService → LicenceService for a restricted, read-only licence-context lookup. The incident service retains its own evidence and case state. It combines evidenceRecorded(i) and contextRecorded(i) before conclusion, as shown by C7 in Part B. A lookup failure cannot be treated as a completed context review.')

new('D Enterprise architecture','foldout')
fig('diagrams/D1_ArchiMate_layers.png','Figure D1. EPA contribution model using the three ArchiMate layers (The Open Group, 2019). The React feature module consumes both workflow APIs; web delivery and the API proxy support it. The exchange XML contains the same elements and relationships.',9.0)

new('D Cross layer explanation')
p('The Business layer describes environmental service responsibilities. The Application layer supplies case-handling functions and their APIs. The Technology layer supports execution and persistence. This separates the environmental purpose from the software that supports it and the infrastructure that runs it (The Open Group, 2019). The diagram uses typed elements; the accompanying Open Group exchange file contains the same model.')
table(['Source element','Relationship','Target element','Meaning'],[
 ('NSW EPA actor','Assignment','Licensing officer and investigator roles','Actor carries the roles'),
 ('Officer/investigator role','Assignment','Licence lifecycle / investigation process','Role performs the behaviour'),
 ('Business process','Realization','EPA-001 / EPA-002 business service','Process implements the externally useful service'),
 ('Workflow application service','Serving','Business process and React EPA module','API supports the process and the user interface'),
 ('Application function','Realization','Workflow application service','Internal behaviour implements the API capability'),
 ('LicenceService / IncidentService component','Assignment','Its application function','Component performs the function'),
 ('Application function','Read/write Access','Owned data object','Function reads and updates its case information'),
 ('Technology service','Serving','Application component','Backend execution/persistence or frontend web delivery/proxy support'),
 ('Local host node','Realization','Technology service','Runtime supplies the execution capability'),
 ('SQLite artifact','Realization','Corresponding data object','Physical file represents the logical data')],[1.62,1.08,1.9,2.27],9.2)
h('Baseline and target')
p('The baseline contains public EPA channels and guidance, including eConnect EPA (NSW EPA, 2026a). The target shown here adds proposed Hub-facing application components around the selected workflows. The local target includes the React EPA module, a Vite delivery/proxy process, two Uvicorn API processes and two SQLite files. It is not a diagram of existing EPA infrastructure.')
h('Integration into the group architecture')
p('The React EPA module can be mounted inside the shared portal. The group owns the surrounding navigation, identity and API routing. LicenceService owns licence data; IncidentService owns incident evidence. Other provider modules use agreed APIs rather than the EPA databases. The shared authentication adapter still needs integration and testing.')
p('The native BPMN files were validated against the OMG BPMN schema and the ArchiMate exchange file against the Open Group 3.1 diagram schema (Object Management Group, 2014; The Open Group, n.d.). Schema validity checks syntax and references accepted by that schema; it does not prove business completeness or regulatory correctness.')

new('E Microservices deployment','landscape')
fig('diagrams/E1_Deployment.png','Figure E1. React and API deployment. Local Vite/Python execution is verified; the alternative Nginx/Compose configuration has not been executed. Each backend owns its persistent database.',5.6)
p('Licence applications and variations share a service because they manage the same licence lifecycle and must update licence versions consistently. Incident reporting is a separate service with a different case lifecycle and evidence. This applies microservice principles of cohesive business boundaries, separate processes and data ownership (Lewis & Fowler, 2014).','Caption')

new('E React frontend and user workflows')
p('The frontend uses React, JavaScript, Vite, React Router and Bootstrap styling. Its EPA feature module contains the pages for both services. Applicants can submit and follow cases, while the officer workspace exposes the checks and decisions allowed by each case state. The API remains responsible for ownership, roles and transition rules.')
table(['Screen or route','What the user can do'],[
 ('/epa','Open the EPA dashboard and select a service'),
 ('/epa/cases','List and filter the cases the account may access'),
 ('/epa/licensing/applications/new','Submit an application with applicant, premises and activity details'),
 ('/epa/licences and the variation form','View issued conditions and request a change to an owned licence'),
 ('/epa/incidents/new','Record a pollution concern and an optional licence reference'),
 ('Case detail and activity history','Supply information, record assessments or triage, review context and inspect the outcome'),
 ('/epa/analytics','View officer activity statistics and export the live CSV/XES logs')],[2.5,4.37],9.4)
h('Frontend to backend requests')
p('The shared client in frontend/src/features/epa/api.js adds the bearer credential, sends JSON and translates API failures into readable messages. Creation requests include an Idempotency-Key; the same key is retained when an unchanged request is retried. State-changing commands include expected_version. An HTTP 409 prompts the user to refresh the case before another action.')
p('Browser requests use /api/v1/epa/licensing or /api/v1/epa/incidents. The development proxy removes the prefix and forwards to port 8001 or 8002. The configured Nginx route performs the equivalent mapping for Compose. The service endpoints in Part C remain relative to their owning backend.')
h('Authentication in the demonstration')
p('The account selector uses published demo tokens for an applicant, a second applicant and an officer. It is not a registration or password-login system, and it does not validate a group JWT. Backend checks enforce the stored owner and permitted role. Replacing the selector with the group login also requires a verified identity dependency in both APIs.')
h('Frontend organisation')
p('EpaRoutes.jsx defines the feature routes; forms.jsx contains the three submission forms; CaseDetail.jsx presents state-dependent actions; pages.jsx contains the dashboard, lists and live activity view. EPA styling is scoped beneath .epa-app. The reusable module receives its identity and API locations through EpaProvider.')

new('E Shared portal and Docker integration')
p('The EPA frontend is designed to become one feature inside the group’s React application. The group’s existing router and login remain the common entry point. EpaRoutes embedded mode omits the development host’s header, sidebar and account selector while retaining the EPA pages and local navigation.')
table(['Integration item','EPA contract or configured default'],[
 ('React mount','/epa/* inside the existing group router'),
 ('Identity supplied to the UI','profile with subject, name, role and token; UI roles client or officer'),
 ('Licensing API prefix','/api/v1/epa/licensing'),
 ('Incident API prefix','/api/v1/epa/incidents'),
 ('Backend identity boundary','Verified principal with stable subject and trusted role in both app factories'),
 ('Service-to-service credential','Backend-only credential for licence context; never included in frontend variables')],[2.05,4.82],9.4)
p('The shared router mounts the existing module as follows. currentUser and accessToken come from the group’s chosen identity system; this snippet assumes authentication has already resolved.')
code("""
<Route path="/epa/*" element={
  <EpaProvider key={currentUser.subject}
    profile={{ ...currentUser, token: accessToken }}
    licensingUrl="/api/v1/epa/licensing"
    incidentsUrl="/api/v1/epa/incidents" basePath="/epa">
    <EpaRoutes embedded />
  </EpaProvider>
} />
""")
p('The group must map its roles to the UI contract and configure both APIs to verify its credentials. Passing a new token to React alone will not replace the demo authentication adapter. Route names, stable user subjects and authorised officer roles must be tested together after integration.')
h('Optional container deployment')
p('compose.yaml defines three containers: licensing, incidents and web. The web container serves compiled React assets through Nginx on localhost port 8080. Each backend has its own database volume. IncidentService calls licensing:8001 on the internal network; browser requests use the public web origin. Docker packages the agreed runtime configuration, but local development can use Python and Node without it.')
p('Docker and Nginx execution remain to be checked on a group machine. Integration also needs deep-link refresh, applicant ownership, officer permissions, internal credential handling and a demonstration that includes another provider’s implemented service.')

new('E Implementation and state transitions')
table(['Service','Source and storage','Implemented behaviour'],[
 ('LicenceService','services/licensing.py; data/licensing.sqlite','P1 applications, requests for information, checks, two assessments, reasoned decisions; P2 ownership, eligibility, assessment and atomic licence version update'),
 ('IncidentService','services/incidents.py; data/incidents.sqlite','P3 report intake, triage routes, evidence, HTTP licence context, prerequisite join and conclusion'),
 ('Shared library','services/domain.py and services/common.py','Pure business transitions, authentication helpers, version checks, idempotency and atomic event persistence')],[1.2,2.6,3.07])
p('Sharing a code library does not share a database or require the two APIs to run in one process. run.py starts two Uvicorn processes. The incident API calls the licensing API over HTTP. The selected design is synchronous orchestration; each service commits its own work and owns its audit log.')
table(['ID','Permitted progression','Important condition'],[
 ('P1','SUBMITTED → ASSESSMENT or AWAITING_INFORMATION; supply returns to SUBMITTED; final APPROVED, REFUSED or INCOMPLETE_CLOSED','Approval requires both positive assessments. Expiry requires the information deadline to have elapsed.'),
 ('P2','SUBMITTED → ASSESSMENT or INVALID; then APPROVED or REFUSED','Approval requires a passed change assessment and the same active licence version as at submission.'),
 ('P3','SUBMITTED → EMERGENCY_GUIDANCE, REFERRED or INVESTIGATION; then ACTION_RECORDED or NO_FURTHER_ACTION','Conclusion requires evidence and context review. Guidance branches do not start an investigation.')],[.55,3.65,2.67])
h('Transport and persistence')
p('Creation endpoints return 201 with a case ID. Protected reads return 200 to the owner or an officer. Missing/invalid credentials return 401; unauthorised roles or ownership return 403; unknown records return 404; invalid transitions or stale versions return 409; malformed inputs return 422. Creation requires an Idempotency-Key, otherwise 400. The licensing-context outage response is 503.')
p('Each successful state mutation, any related licence update and its completed task events are committed together in one SQLite transaction. The expected_version guard prevents a lost case update. A separate licence-version guard prevents one approved variation from silently overwriting another. API event timestamps are real execution times; virtual business durations are generated separately for Part F.')

new('E Code excerpts from the implementation')
h('Require both assessment branches before a decision')
code('''
require(state == 'ASSESSMENT' and 'environment' in d
        and 'suitability' in d,
        'Both assessments must finish before a decision')
approved = values['approved']
require(not approved or (d['environment']['passed']
                        and d['suitability']['passed']),
        'Approval requires both assessments to pass')
case['state'] = 'APPROVED' if approved else 'REFUSED'
events = ['A8', 'A9' if approved else 'A10', 'A11']
''')
p('Source: services/domain.py, selected lines from the P1 decide branch. The HTTP handler also requires a boolean decision and a nonempty rationale. Returning task IDs keeps the API and simulation aligned with the BPMN indexing.')
h('Keep state and audit events in the same transaction')
code('''
with self.connect() as db:
    db.execute('BEGIN IMMEDIATE')
    # Read the stored case and authorise its owner or officer.
    if old['version'] != expected:
        raise HTTPException(409,
                            'Stale version; reload before retrying')
    new, tasks = function(old, db)
    new['version'] = old['version'] + 1
    db.execute('UPDATE cases SET body=? WHERE id=?',
               (json.dumps(new), case_id))
    self.append_events(db, new, tasks, user['subject'])
''')
p('Source: services/common.py, Store.apply, abbreviated to show the transaction and version control. The complete method in the package includes missing-case checks, ownership enforcement and conversion of invalid business transitions into 409 responses.')
h('Compose the incident and licensing services')
code('''
response = httpx.get(
    licence_url + '/internal/licences/' + licence_id + '/context',
    headers={'Authorization': 'Bearer ' + internal_token},
    timeout=2.0, trust_env=False)
# 404 becomes an explicit no-match result.
# Other HTTP/transport failures return 503 without committing I7.
''')
p('Source: services/incidents.py, lookup, with environment lookups abbreviated as licence_url and internal_token. The actual implementation catches transport failures, avoids forwarding owner details and leaves the saved incident available for retry.')

new('E Verification and practical limits')
p('Using FastAPI’s testing facilities (FastAPI, n.d.), the recorded verification completed 19 API tests. Seven JavaScript helper tests, the React build and six browser workflows run individually also passed. A two-server HTTP check completed application → variation → incident and retrieved licence version 2. Evidence is in results/api_v2_tests.txt, results/frontend_browser_tests.txt, results/http_demo.json and docs/VERIFICATION.md. This revision reused those results; it did not rerun the software.')
table(['Verified behaviour','Evidence or implication'],[
 ('Credentials and owner boundaries','Invalid credentials rejected; another client cannot read the owner’s case; officer commands require the officer role'),
 ('Idempotency and stale versions','Identical submission returns the same case; changed payload or stale mutation returns 409 without duplicate completion events'),
 ('Process prerequisites','Premature decision or conclusion is rejected; failed assessment cannot be approved'),
 ('Rework and timer','Missing information can be supplied and rechecked; expiry is rejected before its due time'),
 ('Variation and persistence','Approved changes update the owned licence; saved records survive app recreation'),
 ('Incident routing and lookup failure','Guidance paths skip investigation; failed context lookup preserves state and version'),
 ('React workflows and HTTP composition','Application, variation and incident work through the UI; the incident receives the updated licence version'),
 ('Frontend behaviour','Ownership rejection, stale-update recovery, duplicate-click handling, CSV/XES downloads and mobile/outage behaviour were exercised')],[2.1,4.77])
h('Limits of the current implementation')
p('Bearer tokens are fixed demonstration identities and the servers bind to localhost. Production deployment would require managed identities, transport encryption, secret management and an agreed authorisation model. Database encryption, file uploads, pagination, malware scanning, durable notifications and background scheduling are not implemented. Evidence and supporting documents are represented by text summaries and references.')
p('The 30-day timer is an explicit officer command with a server-side time check. The domain simulation advances a virtual clock to exercise that branch. A production timer worker would need safe retry and scheduling, while preserving the same state guard.')
p('The two-second context timeout is implemented; monthly availability and response-percentile targets were not benchmarked. Local Python/Node execution is the verified route. Compose and Nginx were not executed. The six browser tests ran separately with a custom Chromium binary; one combined run with a standard Playwright installation remains to be checked. Shared portal and identity integration is also pending.')
p('Business decisions are officer inputs in the prototype. The model does not encode the full legislative decision space, exercise delegated authority or submit data to official channels. The real reporting guidance directs immediate danger to 000 and provides EPA contact channels (NSW EPA, 2025c); the prototype receipt points users to those channels.')

new('F Execution and event logs')
p('Part F executes the three processes using a seeded simulation, with virtual task durations. analytics/simulate.py calls the same transition function as the APIs, emits the resulting task IDs, and schedules virtual task and waiting durations. Seed 927 with 100 cases per process produced 300 completed cases and 2,116 events.')
table(['Log field','Meaning'],[
 ('case_id, process','Trace identity and P1/P2/P3 process'),('task_id, activity','Stable task key and readable name'),
 ('start_timestamp, timestamp','Task start and completion with timezone'),('resource, lifecycle','Responsible role; completion lifecycle marker'),
 ('active_hours, wait_hours','Generated task work and preceding queue delay'),('event_index, outcome','Completion order within a case and final case outcome')],[2.35,4.52])
p('Each event records a case ID, task ID, timestamp and resource. XES organises event data into traces and supports standard attributes for names, times, lifecycle transitions and resources (Fluxicon, 2010). In this project, task_id maps to concept:name, completion to time:timestamp and role to org:resource. Additional start times support duration analysis. Cases are sorted by completion time and then event index.')
h('Simulation assumptions')
p('Each non-timer wait is the configured mean multiplied by a uniform factor from 0.7 to 1.3. Active duration uses a factor from 0.75 to 1.25. A5 waits exactly 720 virtual hours after the information request. Both parallel tasks begin scheduling from the same clock; the join resumes at the later completion, not their summed durations.')
table(['Decision or duration','Configured assumption'],[
 ('P1 initial incomplete / timeout if incomplete','35% / 15%'),('P1 second information round if responding','20%'),
 ('P1 assessment failures','Environmental 15%; suitability 8%, sampled independently'),('P2 ineligible / failed assessment if eligible','12% / 20%'),
 ('P3 emergency / referral / EPA route','10% / 20% / 70%'),('P3 substantiated given investigation','60%'),
 ('Main queue means','A4/A6/A7: 48 hours; V4/I7: 24 hours; I6: 18 hours')],[3.3,3.57])
p('These inputs are invented for the experiment and recorded in the source and simulation metadata. They are neither fitted to EPA observations nor estimates of its service times. P2 uses synthetic eligible licences in simulation; real cross-service consistency is exercised separately by the HTTP demonstration.')

new('F Measured simulation outcomes')
table(['ID','Cases','Events','Variants','Mean h','Median h','90th pct h'],[
 (k,v['cases'],v['events'],v['variants'],f"{v['cycle_hours_mean']:.2f}",f"{v['cycle_hours_median']:.2f}",f"{v['cycle_hours_p90']:.2f}") for k,v in metrics.items()],[.65,.6,.68,.78,1.03,1.16,1.97])
p('Case duration is the last completion timestamp minus the first task start. It includes simulated waiting and respects parallel overlap. These are virtual hours, not request latency.')
table(['ID','Terminal outcomes'],[
 ('P1','70 approved; 26 refused; 4 incomplete cases closed'),('P2','60 approved; 26 refused; 14 invalid'),
 ('P3','48 action recorded; 19 no further action; 21 referral guidance; 12 emergency guidance')],[.65,6.22])
fig('results/performance.png','Figure F1. Computed case durations and mean task queue delays from the generated log. The explicit 720-hour information timer is excluded from the right-hand ranking.',3.0)
p('P1 had 36 cases with at least one information request, including seven with a second round. Their mean case duration was 216.44 hours, compared with 85.70 hours for cases without a request. P3’s mean is below its median because short guidance paths are mixed with longer investigations. The observed differences reflect the chosen model and parameters; they are not evidence of causal effects in real EPA operations.')

new('F Discovery and a worked Alpha example')
p('analytics/mine.py uses PM4Py 2.7.23.8 for discovery and conformance checking (Process Intelligence Solutions, n.d.-a, n.d.-b). For each process, the first 80 sorted case IDs form the training set and the remaining 20 are held out. It runs Alpha and Inductive discovery, exports Petri nets as PNML and computes token-replay fitness on the held-out traces. It also exports directly-follows frequencies, footprint tables, variants and task performance.')
p('The Alpha algorithm’s footprint uses x > y for direct succession, x → y when only x > y occurs, x || y when both directions occur, and x # y when neither occurs (van der Aalst et al., 2004). P2 supplies a compact worked example because its three observed variants contain sequence and exclusive choice without loops.')
eq('W₂ = {⟨V1,V2,V3,V8⟩, ⟨V1,V2,V4,V5,V6,V8⟩, ⟨V1,V2,V4,V5,V7,V8⟩}')
eq('T = {V1,V2,V3,V4,V5,V6,V7,V8}; Tᵢ = {V1}; Tₒ = {V8}')
with (ROOT/'results/P2_footprint.csv').open() as f:rows=list(csv.reader(f))
rows=[[x.replace('->','→').replace('<-','←') for x in row] for row in rows]
table(rows[0],rows[1:],[.79]+[.76]*8,9)
p('The maximal Alpha pairs YW are shown below. Each pair creates a place with incoming arcs from the activities in its left set and outgoing arcs to its right set. Activities within either set must be pairwise unrelated (#). Add a source place before V1 and a sink after V8.')
eq('YW = {({V1},{V2}), ({V2},{V3,V4}), ({V4},{V5}),')
eq('          ({V5},{V6,V7}), ({V3,V6,V7},{V8})}')
p('This yields seven places including source and sink, and eight labelled transitions. The discovered structure matches the eligibility choice, approval/refusal choice and common outcome step. P1’s footprint identifies A6 || A7; P3 identifies I6 || I7. Both completion orders appear in each corresponding log.')

new('F The discovered P2 model','landscape')
fig('results/mined_models/P2_alpha.png','Figure F2. The actual PM4Py Alpha net for P2, rendered from the mined graph. The initial place contains a token; V8 leads to the final place.',4.7)
p('The left branch from V2 reaches V3 and then the common pre-V8 place. The other branch reaches V4 and V5, then chooses V6 or V7. No branch requires both approval and refusal. PNML files for Alpha and Inductive models of all three processes are supplied in results/mined_models/.','Caption')

new('F Conformance and interpretation')
table(['ID','Alpha places / transitions','Inductive places / transitions','Held out fitting traces','Reference fitness'],[
 (k,f"{v['discovery']['alpha']['places']} / {v['discovery']['alpha']['transitions']}",f"{v['discovery']['inductive']['places']} / {v['discovery']['inductive']['transitions']}",'20 / 20 for both','1.000') for k,v in metrics.items()],[.7,1.45,1.5,1.65,1.57])
p('Both discovered models scored held-out token-replay log fitness of 1.000 for each process. All 100 traces per process also fitted a separately authored reference Petri net representing the BPMN control flow. The reference models are built in analytics/reference_nets.py; they are not discovered from the training log.')
table(['Negative control','Removed prerequisite','Fitting traces','Token fitness'],[
 ('P1 altered trace','A7 suitability assessment','0 of 1','0.8889'),('P2 altered trace','V4 change assessment','0 of 1','0.8333'),('P3 altered trace','I7 licence context review','0 of 1','0.9000')],[1.5,2.4,1.45,1.52])
p('The negative controls demonstrate that the reference-net replay detects these deliberately missing prerequisites. The nonzero fitness values mean much of each altered trace still fits; none is a fully fitting execution.')
h('What the results establish')
p('The generated traces are consistent with the selected model’s ordering and branches, and the held-out examples are replayable by the discovered models. This is expected for deliberately clean data generated from the design. Fitness alone does not measure precision, model simplicity, predictive performance, statutory compliance or coverage of real-world exceptions. No precision score or production performance claim is made.')
p('Petri-net replay here checks control flow only. It does not evaluate officer authority, case ownership, boolean assessment guards, licence-version conflicts or the 30-day clock. Those selected rules are addressed by API tests and code guards. The 80/20 split is deterministic within one synthetic generator; it is not independent validation against external EPA data.')
h('Useful changes to evaluate next')
p('For the simulated application process, clearer submission prompts and a completeness checklist are plausible ways to reduce information rework. The next experiment should change only the incomplete-case probability and compare cycle-time distributions over several seeds. Assessor capacity and queue policies can be studied separately because the current largest queues are configured around the assessment and information steps.')
p('For incident analysis, compare emergency/referral routes separately from investigations. A single overall mean conceals their different workloads. Predictive delay models would need a suitable independent dataset and evaluation plan; none is trained here. Classic Alpha also assumes sufficiently complete, clean observations and can struggle with short loops and noise (van der Aalst et al., 2004). P1’s modelled rework loop spans three task labels; Inductive discovery supplies an additional comparison, not a proof that all Alpha limitations have been removed.')

new('F Live activity and group event logs')
p('The React Process activity page reads committed API events. It shows case counts, observed task-sequence variants and elapsed times between recorded events. These values describe the local demonstration session. They are separate from the 300 simulated cases and virtual business-hour results analysed in this report; the browser does not run Alpha or Inductive mining.')
table(['Event source','Recorded information','Meaning and use'],[
 ('Seeded simulation','case_id, process, task_id, activity, start/completion timestamps, resource, active_hours and wait_hours','Virtual work and waiting durations; source of the Part F experiment'),
 ('Live API CSV','case_id, process, task_id, activity, timestamp, resource, actor and seq','Actual task-completion events; no simulated work/wait duration fields'),
 ('Live API XES','Trace case ID and process; activity name, task_id, resource, completion timestamp and lifecycle','Export for inspection or a separately prepared process-mining analysis')],[1.35,2.5,3.02],9.3)
h('Prepare a common group log')
p('Retain the original EPA exports and create a separate normalised copy for group analysis. Add provider_id=EPA and service_id=EPA-001 for P1/P2 or EPA-002 for P3. Qualify process and task labels, for example EPA.P1 and EPA.P1.A1. Build a trace key from provider, process and original case_id so another module’s identifiers cannot merge into an EPA case.')
p('The simulated XES uses task IDs as concept:name, while the live XES uses activity names and keeps task_id separately. Choose one common classifier before comparing them. Using the qualified task_id retains the mapping back to the BPMN activity. Preserve timezone information and the original event ordering; use seq to distinguish tied live CSV timestamps.')
p('Different provider cases are separate traces unless an implemented orchestration supplies a shared business-case identifier and a documented relationship. Combining records solely because they belong to the same user would produce a misleading process model. Likewise, do not pool live demonstration timings with the simulator’s virtual durations.')
h('Evidence still required from the group')
p('Each implemented provider needs its own execution evidence and an agreed event mapping. The combined analysis should identify which log was mined, the algorithm and settings, the case count and the resulting model. No event-log merger or cross-provider mining result is claimed by the EPA prototype.')

new('G Contribution evidence and group review')
p('Contribution ratings must describe actual work and be assessed by the other members. The assignment requires the ratings and detailed justification on the group report’s cover page, with each member’s own agreement and signature. This package cannot establish attendance, authorship history or a peer-assessed rating on a person’s behalf.')
table(['Area','Available evidence for review','Actual personal work to confirm'],[
 ('A','Two BSRL specifications and source catalogue','Research, choice of services, requirements review or revisions actually performed'),
 ('B','Three BPMN files; 31 immediate effects; cumulative scenarios','Modelling, semantic reasoning, review or corrections actually performed'),
 ('C and D','SoaML figures; ArchiMate exchange model','Architecture decisions, integration and explanation actually performed'),
 ('E','React module, API source, OpenAPI, tests and HTTP/browser evidence','Implementation, debugging, execution or code review actually performed'),
 ('F','Simulator, 300 cases, logs, miners and metrics','Experiment design, reproduction and interpretation actually performed'),
 ('Group work','Meeting minutes, integration changes and demonstration recording to be supplied by the team','True attendance, agreed actions, reviews and integration contributions')],[.8,2.8,3.27])
h('Personal statement worksheet')
p('Member: Muhammad Nafees Shihab    Student number: __________________')
p('Section summary for the group report: NSW EPA is the allocated provider. The supplied section covers EPA-001 and EPA-002 through BSRL, three annotated BPMN processes, SoaML, ArchiMate, a React module, two microservices and process analytics. Confirm the personal research, implementation, review and integration work below before assigning a contribution rating.')
p('Work I personally completed or reviewed, with file or commit evidence:')
p('____________________________________________________________________')
p('____________________________________________________________________')
p('Changes I made after team review and how I tested or checked them:')
p('____________________________________________________________________')
p('Peer-assessed contribution rating: ______________________________________')
p('Justification agreed by the other members: _______________________________')
p('____________________________________________________________________')
p('Use the subject’s permitted percentage or categorical rating scheme. A percentage contribution rating is not automatically an equal division of the group’s workload. Do not assign 100% or “contributed” without the required peer assessment. Any assistance used should be acknowledged in the way required by the subject.')

new('G Member agreements and signatures')
p('Complete these fields after the group has reviewed the actual contribution report and ratings. Each member must personally agree and provide their own signature. Copy the agreed ratings and justification to the final group cover page.')
for n in range(1,7):
    h('Member '+str(n))
    p('Name: '+('Muhammad Nafees Shihab' if n==1 else '________________________')+'    Rating: __________________')
    p('I agree with my group member contribution report and ratings')
    p('Signature: ______________________________    Date: __________________')
p('These blank agreement fields are intentionally unconfirmed. No member’s signature, consent, attendance or contribution rating has been supplied or inferred.')

new('Reproduction and submission guide')
h('Run the prototype')
p('Install Python 3.12 and Node.js 24, then open a terminal in the package root. The commands below use an activated Python virtual environment. README.md provides the Windows PowerShell and macOS/Linux activation steps.')
code("""
python -m venv .venv
# Activate .venv for your operating system, then:
python -m pip install -r requirements-api.txt
cd frontend
npm ci
cd ..
python run.py --with-frontend
""")
p('Open http://127.0.0.1:5173 and choose an applicant or officer demo account. Create an application, approve a variation, then investigate an incident against that licence. The launcher starts both APIs and Vite; Ctrl+C stops them. To use the configured container route instead, run docker compose up --build and open localhost port 8080; that route still requires group-machine validation.')
h('Reproduce checks and analytics')
code("""
python -m unittest tests.test_api tests.test_frontend_api -v
python -m tests.http_integration
cd frontend
npm test
npm run build
npx playwright install chromium
npm run test:e2e
""")
p('Stop normal development servers before the browser tests, which start isolated services on the same ports. Set the documented PYTHON option if the test launcher needs an explicit virtual-environment executable. Return to the package root for the separate seeded analysis:')
code("""
python -m pip install -r requirements.txt
python -m analytics.simulate --cases 100 --seed 927
python -m analytics.mine
""")
p('Re-running the simulator replaces its generated result files. Keep live exported logs separate. docs/GROUP_INTEGRATION.md describes the React mount, authentication dependency and API proxy rules. The package records local verification and does not establish that the combined group application has passed these checks.')

h('Submission check')
p('The supplied specification sets the final report, executable/source and demonstration recording deadline at 11:00am Wednesday 21 October 2026. It also requires a live demonstration, a recording of no more than ten minutes, and true weekly meeting reports for Weeks 3–12. Check the subject’s current submission channel for any later lecturer update.')
p('The next page lists the remaining integration and submission actions. README.md provides a longer rehearsal outline.')

new('Actions before the group submission')
p('The technical EPA section is ready for personal review and integration. These remaining actions depend on the actual group details and submission arrangements.')
table(['Action','Information or evidence needed'],[
 ('Complete the cover details','Student number, confirmed group ID and the final group cover format'),
 ('Confirm the group catalogue','The other providers and their two services; check that all service IDs are unique'),
 ('Connect the group application','Shared React router, verified identity dependency in both APIs, public API prefixes and another provider’s runnable module'),
 ('Meet the group implementation minimum','At least one additional implemented service from another provider, with its demonstration'),
 ('Review the EPA prototype locally','Run API/frontend tests, build and browser workflows; explain BPMN branches, ownership and state/version checks'),
 ('Complete Part G honestly','Each member’s actual work, peer-agreed rating and personal signature; genuine meeting records'),
 ('Record the combined demonstration','A group recording no longer than ten minutes and preparation for the live demonstration'),
 ('Assemble the final hand-in','One integrated report plus executable/source and the required recording; confirm Moodle instructions')],[2.55,4.32],9.3)
h('Suggested EPA segment of the group demonstration')
p('Allow about two minutes for this provider in the combined video. Show the two services and the P1 information loop first. Run the prepared React workflow, then point out the licence version changing to 2 and the incident retrieving that version. Finish with the P2 discovered model and one analytics finding. Keep the longer eight-minute outline in README.md for rehearsal or a standalone explanation.')
h('Questions to be ready to explain')
p('Why are application and variation inside one business service? Why does an incident need both evidence and licence context? What is the difference between an immediate effect and a cumulative scenario? Which values are simulated? Why does perfect replay fitness not prove real-world correctness? These questions connect the code and diagrams to the underlying modelling concepts.')
p('The final Part G statement should describe the work actually performed and reviewed. A prepared technical package is evidence to inspect; it cannot establish a member’s attendance, effort or agreement by itself.')

new('References')
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def add_reference(author_date, title, publication, url):
    para = p()
    para.paragraph_format.left_indent = Inches(.25)
    para.paragraph_format.first_line_indent = Inches(-.25)
    para.paragraph_format.space_after = Pt(11)
    para.paragraph_format.line_spacing = 1.08
    para.paragraph_format.keep_together = True
    para.add_run(author_date)
    para.add_run(title).italic = True
    para.add_run(publication)
    relation = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), relation)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1155CC'); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '20'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = url; run.append(text)
    link.append(run); para._p.append(link)

references = [('FastAPI. (n.d.). ', 'Testing.', ' FastAPI documentation. Retrieved from ', 'https://fastapi.tiangolo.com/tutorial/testing/'), ('Fluxicon. (2010, September 21). ', 'An introduction to the XES standard.', ' Flux Capacitor. Retrieved from ', 'https://fluxicon.com/blog/2010/09/intro-to-xes/'), ('Ghose, A. K., Lê, L. S., Hoesch-Klohe, K., & Morrison, E. (2011). ', 'The business service representation language: A preliminary report.', ' In M. Cezon & Y. Wolfsthal (Eds.), Towards a Service-Based Internet. ServiceWave 2010 Workshops (pp. 145–152). Springer. ', 'https://doi.org/10.1007/978-3-642-22760-8_16'), ('Hinge, K., Ghose, A. K., & Koliadis, G. (2009). ', 'Process SEER: A tool for semantic effect annotation of business process models.', ' In 2009 IEEE International Enterprise Distributed Object Computing Conference (pp. 54–63). IEEE. ', 'https://doi.org/10.1109/EDOC.2009.24'), ('Lewis, J., & Fowler, M. (2014, March 25). ', 'Microservices.', ' MartinFowler.com. Retrieved from ', 'https://martinfowler.com/articles/microservices.html'), ('NSW Environment Protection Authority. (2020, November 25). ', 'Authorised officers and enforcement officers.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Authorised-Officers-and-Enforcement-Officers'), ('NSW Environment Protection Authority. (2022, October 10). ', 'Compliance audit program.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Licensing/Environment-protection-licences/Compliance-audit-program'), ('NSW Environment Protection Authority. (2024a, August 5). ', 'Financial assurance.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Regulation/Policies-and-guidelines/Financial-assurance'), ('NSW Environment Protection Authority. (2024b, August 5). ', 'Financial capacity.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Regulation/Policies-and-guidelines/Financial-capacity'), ('NSW Environment Protection Authority. (2024c). ', 'Regulatory policy.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Publications/About/2024p4550-Regulatory-Policy'), ('NSW Environment Protection Authority. (2024d, July 3). ', 'What we do.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/About-us/What-we-do'), ('NSW Environment Protection Authority. (2025a, December 10). ', 'EPA customer service charter.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/about-us/epa-customer-service-charter'), ('NSW Environment Protection Authority. (2025b, March). ', 'Prosecution guidelines.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Publications/legislation/22j0619-Prosecution-Guidelines'), ('NSW Environment Protection Authority. (2025c, November 17). ', 'Report pollution.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Reporting-and-incidents/Report-pollution'), ('NSW Environment Protection Authority. (2026a, September 21). ', 'eConnect EPA.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Licensing/econnect-epa'), ('NSW Environment Protection Authority. (2026b, April 28). ', 'Environment protection licences.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Licensing/Environment-protection-licences'), ('NSW Environment Protection Authority. (2026c, April 28). ', 'Guide to licensing.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Licensing/Environment-protection-licences/Guide-to-licensing'), ('NSW Environment Protection Authority. (2026d, April 28). ', 'Licence forms.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Licensing/Environment-protection-licences/Licence-forms'), ('NSW Environment Protection Authority. (2026e, February 26). ', 'Public registers.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Licensing-and-Regulation/Public-registers'), ('NSW Environment Protection Authority. (2026f, April 28). ', 'Reporting pollution.', ' Retrieved from ', 'https://www.epa.nsw.gov.au/Reporting-and-incidents/Report-pollution/contacts-chemical-radiation-pollution/reporting-pollution-overview'), ('Object Management Group. (2012). ', 'Service oriented architecture Modeling Language (SoaML) specification (Version 1.0.1).', ' Retrieved from ', 'https://www.omg.org/spec/SoaML/1.0.1'), ('Object Management Group. (2014). ', 'Business process model and notation (BPMN) specification (Version 2.0.2).', ' Retrieved from ', 'https://www.omg.org/spec/BPMN/2.0.2'), ('Process Intelligence Solutions. (n.d.-a). ', 'Conformance checking [Jupyter notebook].', ' PM4Py. GitHub. Retrieved from ', 'https://github.com/process-intelligence-solutions/pm4py/blob/release/notebooks/4_conformance_checking.ipynb'), ('Process Intelligence Solutions. (n.d.-b). ', 'pm4py.discovery module.', ' PM4Py documentation. Retrieved from ', 'https://processintelligence.solutions/app/pm4py/api/api/pm4py.discovery.html'), ('The Open Group. (n.d.). ', 'ArchiMate model exchange file format for the ArchiMate 3.1/3.2 modeling language.', ' Retrieved from ', 'https://www.opengroup.org/xsd/archimate/'), ('The Open Group. (2019). ', 'ArchiMate 3.1 specification.', ' Retrieved from ', 'https://publications.opengroup.org/c197'), ('van der Aalst, W. M. P., Weijters, A. J. M. M., & Maruster, L. (2004). ', 'Workflow mining: Discovering process models from event logs.', ' IEEE Transactions on Knowledge and Data Engineering, 16(9), 1128–1142. ', 'https://doi.org/10.1109/TKDE.2004.47')]
for reference in references:
    add_reference(*reference)


D.core_properties.title='NSW EPA services for NetZero Hub'
D.core_properties.subject='Individual contribution for Tasks A to G'
D.core_properties.author='Prepared for Muhammad Nafees Shihab'
D.core_properties.keywords='CSCI927 CSCI427 NSW EPA BPMN SoaML ArchiMate process mining'
out=ROOT/'NSW_EPA_Report.docx'
D.save(out)
print(out)
print('Sections:',len(D.sections),'paragraphs:',len(D.paragraphs),'tables:',len(D.tables))
